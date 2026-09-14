from __future__ import annotations

import hashlib
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

import sys

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from server.excel_scope import ExcelScopeError, resolve_excel_scope  # noqa: E402


class ExcelScopeTests(unittest.TestCase):
    def setUp(self) -> None:
        self.directory = tempfile.TemporaryDirectory()
        folder = Path(self.directory.name)
        self.workbook = folder / "baseline.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.title = "目录"
        sheet.cell(1, 1).value = "专题1：实数"
        sheet.cell(2, 2).value = "【大题】"
        sheet.cell(3, 3).value = "实数计算"
        sheet.cell(4, 4).value = "考法1：化简"
        sheet.cell(5, 3).value = "实数应用"
        sheet.cell(6, 2).value = "【微专题】"
        book.save(self.workbook)
        book.close()
        digest = hashlib.sha256(self.workbook.read_bytes()).hexdigest()
        self.config_path = folder / "settings.local.yaml"
        self.settings = {"directory_workbook": {"path": "baseline.xlsx", "sheet": "目录"}}
        self.taxonomy = {
            "source_sheet": "目录", "source_workbook_sha256": digest,
            "topics": [{"id": "topic", "title": "专题1：实数", "source_row": 1, "level2": [{
                "id": "large", "title": "【大题】", "source_row": 2, "level3": [
                    {"id": "calculation", "title": "实数计算", "source_row": 3, "level4": []},
                    {"id": "application", "title": "实数应用", "source_row": 5, "level4": []},
                ],
            }] }],
        }

    def tearDown(self) -> None:
        self.directory.cleanup()

    def test_level2_scope_excludes_the_container_and_stops_at_next_level2(self) -> None:
        scope = resolve_excel_scope(self.settings, self.config_path, self.taxonomy, {
            "level": 2, "topic_id": "topic", "level2_id": "large", "level3_id": None,
        })
        self.assertEqual(scope["container_row"], 2)
        self.assertEqual((scope["replace_start_row"], scope["replace_end_row"]), (3, 5))

    def test_level3_scope_includes_only_its_level4_children(self) -> None:
        scope = resolve_excel_scope(self.settings, self.config_path, self.taxonomy, {
            "level": 3, "topic_id": "topic", "level2_id": "large", "level3_id": "calculation",
        })
        self.assertEqual(scope["container_row"], 3)
        self.assertEqual((scope["replace_start_row"], scope["replace_end_row"]), (4, 4))

    def test_scope_refuses_a_baseline_that_changed_since_taxonomy_export(self) -> None:
        self.taxonomy["source_workbook_sha256"] = "0" * 64
        with self.assertRaises(ExcelScopeError):
            resolve_excel_scope(self.settings, self.config_path, self.taxonomy, {
                "level": 2, "topic_id": "topic", "level2_id": "large", "level3_id": None,
            })


if __name__ == "__main__":
    unittest.main()
