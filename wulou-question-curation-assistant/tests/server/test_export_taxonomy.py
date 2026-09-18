from __future__ import annotations

import importlib.util
import re
import tempfile
import unittest
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook


ROOT = Path(__file__).resolve().parents[2]
MODULE_PATH = ROOT / "scripts" / "export-taxonomy.py"
SPEC = importlib.util.spec_from_file_location("export_taxonomy_script", MODULE_PATH)
if SPEC is None or SPEC.loader is None:
    raise RuntimeError("无法加载目录导出脚本")
MODULE = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(MODULE)

# ⚠️ openpyxl 3.1.5 写的是 `<pageMargins ... footer="0.5" />`——`/` 前面有空格。
#    以前这里按没有空格的串做字面替换，永远匹配不上，于是"空边距"用例其实一直在读
#    一个完全正常的文件、根本没走到兼容副本分支。改成按元素整体替换。
_PAGE_MARGINS = re.compile(rb"<pageMargins[^>]*/>")
_BLANK_PAGE_MARGINS = (
    b'<pageMargins left="" right="" top="" bottom="" header="0.3" footer="0.3"/>'
)


def blank_out_page_margins(workbook_path: Path) -> bool:
    """把工作簿里的页边距改写成空值，模拟题湖导出器的写法。返回是否真的改到了。"""
    rewritten = workbook_path.with_name("rewritten-" + workbook_path.name)
    changed = False
    with ZipFile(workbook_path) as source, ZipFile(rewritten, "w", ZIP_DEFLATED) as destination:
        for member in source.infolist():
            content = source.read(member.filename)
            if member.filename == "xl/worksheets/sheet1.xml":
                content, count = _PAGE_MARGINS.subn(_BLANK_PAGE_MARGINS, content)
                changed = changed or count > 0
            destination.writestr(member, content)
    rewritten.replace(workbook_path)
    return changed


class ExportTaxonomyTests(unittest.TestCase):
    def test_uses_column_a_topic_and_does_not_promote_previous_level2(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            workbook_path = Path(directory) / "taxonomy.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.title = "目录"
            sheet.cell(1, 1).value = "专题12：锐角三角函数"
            sheet.cell(2, 2).value = "12.4 一般角的三角函数值"
            sheet.cell(3, 3).value = "一般角三角函数"
            sheet.cell(4, 2).value = "【大题】"
            sheet.cell(5, 3).value = "实数综合计算（含三角比）"
            sheet.cell(5, 5).value = "TRIG-001"
            sheet.cell(5, 14).value = "须使用三角比完成综合计算。"
            sheet.cell(6, 4).value = "考法1：三角比综合计算"
            sheet.cell(6, 14).value = "须由三角比建立边长关系后求值。"
            sheet.cell(7, 2).value = "【微专题】三角函数"
            sheet.cell(8, 3).value = "不得进入大题"
            book.save(workbook_path)
            book.close()

            exported = MODULE.export_taxonomy(workbook_path, "目录")

        self.assertTrue(exported["taxonomy_version"].startswith("excel-v4-"))
        self.assertEqual(len(exported["topics"]), 1)
        topic = exported["topics"][0]
        self.assertEqual(topic["title"], "专题12：锐角三角函数")
        self.assertEqual(topic["order"], 12)
        self.assertEqual(topic["source_row"], 1)
        self.assertEqual(topic["level2"][0]["title"], "【大题】")
        self.assertEqual(topic["level2"][0]["source_row"], 4)
        self.assertEqual(
            topic["level2"][0]["level3"][0]["title"],
            "实数综合计算（含三角比）",
        )
        self.assertEqual(topic["level2"][0]["level3"][0]["source_row"], 5)
        self.assertEqual(topic["level2"][0]["level3"][0]["classification_basis"], "须使用三角比完成综合计算。")
        self.assertEqual(topic["level2"][0]["level3"][0]["level4"][0]["classification_basis"], "须由三角比建立边长关系后求值。")
        self.assertEqual(exported["source_sheet"], "目录")
        self.assertEqual(len(exported["source_workbook_sha256"]), 64)
        self.assertNotIn("12.4 一般角的三角函数值", str(exported["topics"]))
        self.assertNotIn("不得进入大题", str(exported["topics"]))

    def test_separates_evidence_ids_from_model_classification_basis(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            workbook_path = Path(directory) / "taxonomy.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.title = "目录"
            sheet.cell(1, 1).value = "专题2：代数式"
            sheet.cell(2, 2).value = "【大题】"
            sheet.cell(3, 3).value = "分式化简与求值"
            sheet.cell(4, 4).value = "考法1：按定义域筛选代入值"
            sheet.cell(4, 14).value = "须依据原式限制筛出可代入值；现有证据题目ID：CS2025ABC_001、CS2025DEF-002"
            book.save(workbook_path)
            book.close()

            exported = MODULE.export_taxonomy(workbook_path, "目录")

        level4 = exported["topics"][0]["level2"][0]["level3"][0]["level4"][0]
        self.assertEqual(level4["classification_basis"], "须依据原式限制筛出可代入值")
        self.assertEqual(level4["audit_evidence_exercise_ids"], ["CS2025ABC_001", "CS2025DEF-002"])

    def test_reads_workbook_with_blank_page_margin_via_temporary_compatible_copy(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            workbook_path = Path(directory) / "taxonomy.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.title = "目录"
            sheet.cell(1, 1).value = "专题1：实数"
            sheet.cell(2, 2).value = "【大题】"
            sheet.cell(3, 3).value = "实数运算"
            book.save(workbook_path)
            book.close()

            self.assertTrue(blank_out_page_margins(workbook_path), "空边距没写进去，用例会失去意义")

            exported = MODULE.export_taxonomy(workbook_path, "目录")

            # ⚠️ 兼容副本分支不能把原工作簿的句柄漏着：Windows 上会把文件锁死，
            #    工作簿随后改不了名、删不掉、Excel 里也打不开。
            workbook_path.unlink()
            self.assertFalse(workbook_path.exists())

        self.assertEqual(exported["topics"][0]["level2"][0]["level3"][0]["title"], "实数运算")

    def test_compatible_copy_also_covers_plain_value_error_from_openpyxl(self) -> None:
        """兼容副本分支必须同时捕 TypeError 和 ValueError。

        openpyxl 3.1.5 把 `float('')` 的 ValueError 包成 TypeError 抛出来，
        所以历史上只写 `except TypeError` 也能跑；但那是它的内部包装行为，
        换版本改成直接抛 ValueError 就会让服务端启动直接崩。这条护栏钉住两者都捕。
        """
        from unittest import mock

        from server import taxonomy_export as export_module

        with tempfile.TemporaryDirectory() as directory:
            workbook_path = Path(directory) / "taxonomy.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.title = "目录"
            sheet.cell(1, 1).value = "专题1：实数"
            sheet.cell(2, 2).value = "【大题】"
            sheet.cell(3, 3).value = "实数运算"
            book.save(workbook_path)
            book.close()
            self.assertTrue(blank_out_page_margins(workbook_path))

            real_load_workbook = export_module.load_workbook
            calls = {"count": 0}

            def raise_value_error_once(path, *args, **kwargs):
                calls["count"] += 1
                if calls["count"] == 1:
                    raise ValueError("could not convert string to float: ''")
                return real_load_workbook(path, *args, **kwargs)

            with mock.patch.object(export_module, "load_workbook", raise_value_error_once):
                exported = MODULE.export_taxonomy(workbook_path, "目录")

        self.assertEqual(calls["count"], 2, "应当先失败一次，再读剥过边距的临时副本")
        self.assertEqual(exported["topics"][0]["level2"][0]["level3"][0]["title"], "实数运算")

    def test_unreadable_workbook_still_raises_when_no_margin_repair_applies(self) -> None:
        """捕得宽不等于吞异常：不是空边距引起的问题，必须原样抛出。"""
        from unittest import mock

        from server import taxonomy_export as export_module

        with tempfile.TemporaryDirectory() as directory:
            workbook_path = Path(directory) / "taxonomy.xlsx"
            book = Workbook()
            book.active.title = "目录"
            book.save(workbook_path)
            book.close()

            # 文件本身没有空页边距 ⇒ 兼容副本不做任何改动 ⇒ 原异常必须照旧抛出。
            with mock.patch.object(
                export_module, "load_workbook", side_effect=ValueError("unrelated failure")
            ):
                with self.assertRaises(ValueError):
                    export_module.export_taxonomy(workbook_path, "目录")


if __name__ == "__main__":
    unittest.main()
