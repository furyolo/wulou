"""目录来源（设置面板可选、可固定）的回归测试。

背景：工作簿挪了文件夹，配置里的旧路径就解析不到，服务连启动都撑不住，
用户也就没机会在前端把路径改回来。这里把「可发现、可切换、坏了也不倒」钉住。
"""

from __future__ import annotations

import http.client
import json
import shutil
import tempfile
import threading
import unittest
from pathlib import Path

import yaml
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[2]

from server.main import CurationServer, ServiceState  # noqa: E402
from server.directory_picker import DirectoryPickerError, pick_path  # noqa: E402
from server.taxonomy_sync import (  # noqa: E402
    AmbiguousWorkbookFolder,
    EmptyWorkbookFolder,
    TaxonomySyncError,
    discover_workbook_families,
    resolve_workbook_anchor,
    summarize_directory_sources,
)


def _write_workbook(path: Path, level3_title: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    sheet = book.active
    sheet.title = "目录"
    sheet.cell(1, 1).value = "专题1：实数"
    sheet.cell(2, 2).value = "【大题】"
    sheet.cell(3, 3).value = level3_title
    book.save(path)
    book.close()


def _settings_text(root: Path, workbook_path: Path, taxonomy_path: Path) -> dict:
    return {
        "host": "127.0.0.1",
        "port": 0,
        "taxonomy_path": str(taxonomy_path),
        "rules_path": str(ROOT / "config" / "classification-rules.yaml"),
        "cache_path": str(root / "cache.sqlite3"),
        # 候选范围钉在临时目录里，避免跟着开发机上的真实素材飘。
        "directory_workbook": {"path": str(workbook_path), "sheet": "目录", "search_roots": [str(root)]},
    }


class ResolveAnchorTests(unittest.TestCase):
    def test_accepts_versioned_workbook_even_if_that_exact_file_is_gone(self) -> None:
        # 锚点只是「哪个文件夹、哪个命名前缀」；用户选的旧版本被删了也应该还能用。
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            anchor = root / "导出目录" / "★目录,ID-3777-2026-09-01 v1.xlsx"
            anchor.parent.mkdir(parents=True)
            self.assertEqual(resolve_workbook_anchor(str(anchor)), anchor.resolve())

    def test_rejects_paths_that_break_automatic_version_tracking(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / "导出目录").mkdir()
            for raw in ("", "   ", "relative/目录 2026-09-01 v1.xlsx"):
                with self.subTest(raw=raw):
                    with self.assertRaises(TaxonomySyncError):
                        resolve_workbook_anchor(raw)
            with self.assertRaises(TaxonomySyncError):
                resolve_workbook_anchor(str(root / "导出目录" / "报表.xlsx"))
            with self.assertRaises(TaxonomySyncError):
                resolve_workbook_anchor(str(root / "导出目录" / "~$目录 2026-09-01 v1.xlsx"))
            with self.assertRaises(TaxonomySyncError):
                resolve_workbook_anchor(str(root / "不存在的文件夹" / "目录,ID-3777-2026-09-01 v1.xlsx"))

    def test_accepts_a_folder_and_uses_the_only_family_inside_it(self) -> None:
        # 前端的选择框只能选文件夹或选文件，两条都必须认；指文件夹时用里面唯一那一族的最新版本。
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            folder = root / "导出目录"
            _write_workbook(folder / "★目录,ID-3777-2026-09-05 v5.xlsx", "旧")
            newest = folder / "★目录,ID-3777-2026-09-18 v2.xlsx"
            _write_workbook(newest, "新")
            # 子目录里的不算数：用户指的是这一层。
            _write_workbook(folder / "子目录" / "★目录,ID-3777-2027-01-01 v1.xlsx", "更深")
            self.assertEqual(resolve_workbook_anchor(str(folder)), newest.resolve())

    def test_folder_with_several_families_reports_them_instead_of_picking_one(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory) / "两族"
            _write_workbook(folder / "★目录,ID-3777-2026-09-10 v1.xlsx", "甲")
            _write_workbook(folder / "别的族,ID-9-2026-01-01 v1.xlsx", "乙")
            with self.assertRaises(AmbiguousWorkbookFolder) as caught:
                resolve_workbook_anchor(str(folder))
            self.assertEqual(len(caught.exception.families), 2)
            self.assertIn("别的族", str(caught.exception))

    def test_folder_without_a_versioned_workbook_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory) / "空的"
            folder.mkdir()
            (folder / "报表.xlsx").write_bytes(b"not-a-workbook")
            with self.assertRaises(EmptyWorkbookFolder):
                resolve_workbook_anchor(str(folder))


class DiscoverFamiliesTests(unittest.TestCase):
    def test_groups_by_folder_and_prefix_and_keeps_only_latest(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            _write_workbook(root / "整理后" / "★目录,ID-3777-2026-09-05 v5.xlsx", "旧")
            _write_workbook(root / "整理后" / "★目录,ID-3777-2026-09-18 v2.xlsx", "新")
            _write_workbook(root / "整理后" / "~$★目录,ID-3777-2026-09-30 v9.xlsx", "临时")
            _write_workbook(root / "整理后" / "别的族,ID-9-2026-01-01 v1.xlsx", "别族")
            _write_workbook(root / "整理后" / "没有版本号.xlsx", "不合规")
            _write_workbook(root / "_归档" / "★目录,ID-3777-2026-12-31 v1.xlsx", "归档")

            families = discover_workbook_families([root])

            keys = sorted((family.prefix, family.revision_date.isoformat(), family.revision) for family in families)
            self.assertEqual(keys, [
                ("★目录,ID-3777-", "2026-09-18", 2),
                ("别的族,ID-9-", "2026-01-01", 1),
            ])
            main = next(family for family in families if family.prefix == "★目录,ID-3777-")
            self.assertEqual(main.version_count, 2)
            self.assertEqual(main.latest.name, "★目录,ID-3777-2026-09-18 v2.xlsx")
            self.assertEqual(main.folder, (root / "整理后").resolve())

    def test_missing_roots_and_depth_limit_do_not_explode(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            _write_workbook(root / "一层" / "二层" / "三层" / "深目录,ID-1-2026-01-01 v1.xlsx", "深")
            self.assertEqual(discover_workbook_families([root / "不存在"]), [])
            self.assertEqual(discover_workbook_families([root], maximum_depth=1), [])
            self.assertEqual(len(discover_workbook_families([root], maximum_depth=3)), 1)


class SummaryTests(unittest.TestCase):
    def test_reports_active_anchor_and_the_reason_when_the_path_is_dead(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            live = root / "导出目录" / "★目录,ID-3777-2026-09-18 v2.xlsx"
            _write_workbook(live, "新")
            _write_workbook(root / "导出目录" / "★目录,ID-3777-2026-09-18 v1.xlsx", "旧")
            dead_anchor = root / "老位置" / "★目录,ID-3777-2026-09-18 v1.xlsx"

            broken = summarize_directory_sources(dead_anchor, [root])
            self.assertFalse(broken["anchor_valid"])
            self.assertEqual(broken["active"], "")
            self.assertIn("未找到目录工作簿", broken["error"])
            self.assertEqual([item["is_active"] for item in broken["candidates"]], [False])

            healthy = summarize_directory_sources(root / "导出目录" / "★目录,ID-3777-2026-09-01 v1.xlsx", [root])
            self.assertTrue(healthy["anchor_valid"])
            self.assertEqual(healthy["active"], str(live.resolve()))
            self.assertEqual([item["is_active"] for item in healthy["candidates"]], [True])
            self.assertEqual([item["is_anchor"] for item in healthy["candidates"]], [True])

    def test_no_configured_path_reports_empty_instead_of_raising(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            summary = summarize_directory_sources(None, [Path(directory)])
            self.assertEqual(summary["anchor"], "")
            self.assertEqual(summary["active"], "")
            self.assertFalse(summary["anchor_valid"])
            self.assertEqual(summary["candidates"], [])


class UpdateDirectorySourceTests(unittest.TestCase):
    def _state(self, root: Path, workbook_path: Path) -> ServiceState:
        settings_path = root / "settings.local.yaml"
        settings_path.write_text(yaml.safe_dump(
            _settings_text(root, workbook_path, root / "taxonomy.yaml"), allow_unicode=True,
        ), encoding="utf-8")
        return ServiceState(settings_path)

    def test_switching_persists_the_anchor_and_reloads_the_snapshot(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            old = root / "老位置" / "目录,ID-3777-2026-09-01 v1.xlsx"
            _write_workbook(old, "旧目录")
            moved = root / "导出目录" / "目录,ID-3777-2026-09-18 v1.xlsx"
            newest = root / "导出目录" / "目录,ID-3777-2026-09-18 v2.xlsx"
            _write_workbook(moved, "中间版本")
            _write_workbook(newest, "新目录")

            state = self._state(root, old)
            try:
                result = state.update_directory_source({"path": str(moved)})

                self.assertEqual(result["active"], str(newest.resolve()))
                self.assertTrue(result["anchor_valid"])
                self.assertIn("新目录", str(state.taxonomy.raw))
            finally:
                state.close()

            saved = yaml.safe_load((root / "settings.local.yaml").read_text(encoding="utf-8"))
            # 落盘的是用户选的锚点，不是解析结果，否则下次没法再跟着自动发现走。
            self.assertEqual(saved["directory_workbook"]["path"], str(moved.resolve()))
            self.assertEqual(saved["directory_workbook"]["sheet"], "目录")

            # 重启后仍然指向同一族的最新版本。
            restarted = ServiceState(root / "settings.local.yaml")
            try:
                self.assertEqual(restarted.taxonomy_sync_status, "up_to_date")
                self.assertEqual(restarted.active_directory_workbook(), str(newest.resolve()))
                self.assertIn("新目录", str(restarted.taxonomy.raw))
            finally:
                restarted.close()

    def test_rejected_input_leaves_the_config_untouched(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            old = root / "老位置" / "目录,ID-3777-2026-09-01 v1.xlsx"
            _write_workbook(old, "旧目录")
            state = self._state(root, old)
            try:
                before = (root / "settings.local.yaml").read_text(encoding="utf-8")
                for bad in ("", "不是绝对路径.xlsx", str(root / "缺文件夹" / "目录,ID-1-2026-01-01 v1.xlsx")):
                    with self.subTest(bad=bad):
                        with self.assertRaises(TaxonomySyncError):
                            state.update_directory_source({"path": bad})
                self.assertEqual((root / "settings.local.yaml").read_text(encoding="utf-8"), before)
                self.assertIn("旧目录", str(state.taxonomy.raw))
            finally:
                state.close()

    def test_broken_source_keeps_service_startable_on_the_last_good_snapshot(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            workbook = root / "老位置" / "目录,ID-3777-2026-09-01 v1.xlsx"
            _write_workbook(workbook, "旧目录")
            state = self._state(root, workbook)
            state.close()

            # 复刻真实事故：工作簿被整批挪进子目录，配置里那个文件夹里已经一个都不剩。
            shutil.move(str(root / "老位置"), str(root / "导出目录" / "老位置"))

            revived = ServiceState(root / "settings.local.yaml")
            try:
                self.assertEqual(revived.taxonomy_sync_status, "failed")
                self.assertIn("未找到目录工作簿", revived.taxonomy_sync_error)
                self.assertEqual(revived.active_directory_workbook(), "")
                # 关键点：服务活着，还带着上一份能用的目录，前端才有机会把路径改回来。
                self.assertIn("旧目录", str(revived.taxonomy.raw))

                # 用户在前端把目录来源指到新位置之后，一切照常。
                fixed = revived.update_directory_source({"path": str(root / "导出目录" / "老位置" / workbook.name)})
                self.assertTrue(fixed["anchor_valid"])
                self.assertEqual(revived.taxonomy_sync_error, "")
            finally:
                revived.close()

    def test_broken_source_without_any_snapshot_still_refuses_to_start(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            with self.assertRaises(TaxonomySyncError):
                self._state(root, root / "老位置" / "目录,ID-3777-2026-09-01 v1.xlsx")


class DirectorySourceHttpTests(unittest.TestCase):
    """HTTP 面：前端只认这几个端点，路由注册漏一个就等于功能没上线。"""

    def setUp(self) -> None:
        self.tempdir = tempfile.TemporaryDirectory()
        self.root = Path(self.tempdir.name)
        workbook = self.root / "老位置" / "目录,ID-3777-2026-09-01 v1.xlsx"
        _write_workbook(workbook, "旧目录")
        settings_path = self.root / "settings.local.yaml"
        settings_path.write_text(yaml.safe_dump(
            _settings_text(self.root, workbook, self.root / "taxonomy.yaml"), allow_unicode=True,
        ), encoding="utf-8")
        # 真的弹窗没法在测试里跑，所以注入一个「替身窗口」：走真实的忙/超时逻辑，
        # 只在最后一步由替身返回文件夹路径。
        self.picked: list[str] = []
        self.picked_path = str(self.root / "导出目录")

        def opener(initial_dir: str) -> str:
            self.picked.append(initial_dir)
            return self.picked_path

        self.state = ServiceState(
            settings_path,
            directory_picker=lambda *, initial_dir="": pick_path(initial_dir=initial_dir, opener=opener),
        )
        self.server = CurationServer(("127.0.0.1", 0), self.state)
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()

    def tearDown(self) -> None:
        self.server.shutdown()
        self.server.server_close()
        self.state.close()
        self.thread.join(timeout=2)
        self.tempdir.cleanup()

    def request(self, method: str, path: str, body: dict | None = None):
        connection = http.client.HTTPConnection("127.0.0.1", self.server.server_port, timeout=5)
        raw_body = json.dumps(body, ensure_ascii=False).encode("utf-8") if body is not None else None
        connection.request(method, path, body=raw_body, headers={"Content-Type": "application/json"})
        response = connection.getresponse()
        payload = json.loads(response.read().decode("utf-8"))
        connection.close()
        return response.status, payload

    def test_directory_workbooks_endpoint_lists_and_health_reports_status(self) -> None:
        status, payload = self.request("GET", "/api/v1/directory-workbooks")
        self.assertEqual(status, 200)
        self.assertTrue(payload["anchor_valid"])
        self.assertEqual(payload["sheet"], "目录")
        self.assertEqual(payload["taxonomy_version"], self.state.taxonomy.version)
        self.assertEqual([item["file_name"] for item in payload["candidates"]], ["目录,ID-3777-2026-09-01 v1.xlsx"])
        self.assertTrue(payload["candidates"][0]["is_active"])

        status, health = self.request("GET", "/health")
        self.assertEqual(status, 200)
        self.assertEqual(health["taxonomy_sync_status"], "up_to_date")
        self.assertEqual(health["taxonomy_sync_error"], "")
        self.assertEqual(health["directory_workbook_active"], payload["active"])

    def test_switching_through_http_updates_the_running_service(self) -> None:
        moved = self.root / "导出目录" / "目录,ID-3777-2026-09-18 v1.xlsx"
        _write_workbook(moved, "新目录")
        status, payload = self.request("POST", "/api/v1/settings/directory-workbook", {"path": str(moved)})
        self.assertEqual(status, 200)
        self.assertEqual(payload["active"], str(moved.resolve()))
        self.assertIn("新目录", str(self.state.taxonomy.raw))
        # 目录变了，后续请求必须看到新版本。
        status, taxonomy = self.request("GET", "/api/v1/taxonomy")
        self.assertEqual(status, 200)
        self.assertIn("新目录", json.dumps(taxonomy, ensure_ascii=False))

    def test_invalid_path_returns_400_with_a_readable_reason(self) -> None:
        status, payload = self.request("POST", "/api/v1/settings/directory-workbook", {"path": "不是绝对路径.xlsx"})
        self.assertEqual(status, 400)
        self.assertIn("完整路径", payload["message"])

    def test_picker_endpoint_returns_the_folder_chosen_on_this_machine(self) -> None:
        status, payload = self.request("POST", "/api/v1/settings/directory-picker", {})
        self.assertEqual(status, 200)
        self.assertEqual(payload["path"], self.picked_path)
        self.assertFalse(payload["cancelled"])
        # 开框位置默认落在当前生效工作簿的文件夹，用户换版本少点几层。
        self.assertEqual(self.picked, [str((self.root / "老位置").resolve())])

    def test_picker_endpoint_reports_a_dialog_that_cannot_open(self) -> None:
        # 本机没有图形会话时不能假装成功，也不能把它算成用户取消。
        def broken(*, initial_dir: str = "") -> dict:
            raise DirectoryPickerError("打不开本机选择窗口：no display name")

        self.state.directory_picker = broken
        status, payload = self.request("POST", "/api/v1/settings/directory-picker", {})
        self.assertEqual(status, 400)
        self.assertEqual(payload["error"], "picker_error")
        self.assertIn("打不开本机选择窗口", payload["message"])

    def test_folder_with_a_single_family_is_accepted_over_http(self) -> None:
        folder = self.root / "整理后"
        _write_workbook(folder / "★目录,ID-3777-2026-09-10 v1.xlsx", "第一版")
        newest = folder / "★目录,ID-3777-2026-09-18 v2.xlsx"
        _write_workbook(newest, "第二版")

        status, payload = self.request("POST", "/api/v1/settings/directory-workbook", {"path": str(folder)})

        self.assertEqual(status, 200)
        self.assertEqual(payload["active"], str(newest.resolve()))
        self.assertTrue(payload["anchor_valid"])
        self.assertIn("第二版", str(self.state.taxonomy.raw))
        saved = yaml.safe_load((self.root / "settings.local.yaml").read_text(encoding="utf-8"))
        self.assertEqual(saved["directory_workbook"]["path"], str(newest.resolve()))

    def test_folder_with_several_families_returns_candidates_instead_of_guessing(self) -> None:
        folder = self.root / "两个族"
        _write_workbook(folder / "★目录,ID-3777-2026-09-10 v1.xlsx", "甲")
        _write_workbook(folder / "别的族,ID-9-2026-01-01 v1.xlsx", "乙")
        before = (self.root / "settings.local.yaml").read_text(encoding="utf-8")

        status, payload = self.request("POST", "/api/v1/settings/directory-workbook", {"path": str(folder)})

        self.assertEqual(status, 409)
        self.assertEqual(payload["error"], "ambiguous_folder")
        self.assertEqual(
            sorted(item["file_name"] for item in payload["families"]),
            ["★目录,ID-3777-2026-09-10 v1.xlsx", "别的族,ID-9-2026-01-01 v1.xlsx"],
        )
        # 猜一组等于静默换掉整套目录：配置和当前目录都必须一个字都没动。
        self.assertEqual((self.root / "settings.local.yaml").read_text(encoding="utf-8"), before)
        self.assertIn("旧目录", str(self.state.taxonomy.raw))

    def test_folder_without_a_versioned_workbook_says_why(self) -> None:
        folder = self.root / "空的"
        folder.mkdir()
        (folder / "报表.xlsx").write_bytes(b"not-a-workbook")

        status, payload = self.request("POST", "/api/v1/settings/directory-workbook", {"path": str(folder)})

        self.assertEqual(status, 409)
        self.assertEqual(payload["error"], "empty_folder")
        self.assertIn("YYYY-MM-DD vN.xlsx", payload["message"])


if __name__ == "__main__":
    unittest.main()
