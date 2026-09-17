from __future__ import annotations

import sys
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from server.directory_refactor import prepare_refactor_context  # noqa: E402
from server.taxonomy import Taxonomy  # noqa: E402
from server.excel_sync import sha256, write_approved_refactor  # noqa: E402


class DirectoryHandoffTests(unittest.TestCase):
    """交接输入只负责确定范围与缺口，不生成方案、不写 Excel。"""

    def setUp(self) -> None:
        self.taxonomy = Taxonomy.from_file(ROOT / "config" / "taxonomy.example.yaml")
        self.questions = [{"exercise_id": str(index)} for index in range(1, 7)]
        self.context = prepare_refactor_context({
            "focus": {
                "level": 3,
                "topic_id": "topic-01-real-numbers",
                "level2_id": "topic-01-large",
                "level3_id": "real-number-calculation",
            },
            "questions": self.questions,
        }, self.taxonomy)

    def test_partial_collection_keeps_successful_questions_and_records_gap(self) -> None:
        context = prepare_refactor_context({
            "focus": {
                "level": 3,
                "topic_id": "topic-01-real-numbers",
                "level2_id": "topic-01-large",
                "level3_id": "real-number-calculation",
            },
            "questions": self.questions,
            "collection": {
                "discovered_question_count": 7,
                "collected_question_count": 6,
                "failed_exercise_ids": ["7"],
                "failed_page_urls": [],
            },
        }, self.taxonomy)
        self.assertEqual(context["collection"]["failed_exercise_ids"], ["7"])
        self.assertEqual(len(context["questions"]), 6)

    def test_stratified_sample_is_carried_into_the_handoff(self) -> None:
        """抽样只改变证据门槛；交接清单必须如实带上抽样模式与原始题量。"""
        context = prepare_refactor_context({
            "focus": {
                "level": 3, "topic_id": "topic-01-real-numbers",
                "level2_id": "topic-01-large", "level3_id": "real-number-calculation",
            },
            "questions": self.questions,
            "collection": {
                "discovered_question_count": 20, "collected_question_count": 6,
                "failed_exercise_ids": [], "failed_page_urls": [],
                "sampling": {"mode": "stratified_page", "source_question_count": 20},
            },
        }, self.taxonomy)
        self.assertEqual(context["collection"]["sampling"]["mode"], "stratified_page")
        self.assertEqual(context["collection"]["sampling"]["source_question_count"], 20)

    def test_handoff_references_the_whole_topic_directory_tree(self) -> None:
        reference = self.context["reference_directory_tree"]
        self.assertGreaterEqual(len(reference), 1)
        self.assertIn("topic-01-large", [level2["id"] for level2 in reference])
        self.assertTrue(all("level3" in level2 for level2 in reference))

    def test_excel_refactor_preserves_values_outside_replaced_range(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            baseline = folder / "baseline.xlsx"
            output = folder / "output.xlsx"
            plan_path = folder / "plan.json"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "目录"
            sheet.cell(1, 1).value = "前缀"
            sheet.cell(2, 3).value = "旧三级"
            sheet.cell(3, 4).value = "旧四级"
            sheet.cell(4, 1).value = "后缀"
            workbook.save(baseline)
            plan_path.write_text(json.dumps({
                "status": "approved", "baseline_sha256": sha256(baseline), "sheet": "目录",
                "replace_start_row": 2, "replace_end_row": 3,
                "rows": [
                    {"level": 3, "title": "新三级", "question_count": 6},
                    {"level": 4, "title": "考法1：新四级", "question_count": 6},
                ],
            }, ensure_ascii=False), encoding="utf-8")
            write_approved_refactor(baseline, output, plan_path)
            result = load_workbook(output, data_only=False)["目录"]
            self.assertEqual(result.cell(1, 1).value, "前缀")
            self.assertEqual(result.cell(2, 3).value, "新三级")
            self.assertEqual(result.cell(3, 4).value, "考法1：新四级")
            self.assertEqual(result.cell(4, 1).value, "后缀")


if __name__ == "__main__":
    unittest.main()
