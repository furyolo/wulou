from __future__ import annotations

import json
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

import server.excel_sync as excel_sync  # noqa: E402
from server.excel_sync import (  # noqa: E402
    inspect_workbook,
    sha256,
    write_approved_plan,
    write_approved_refactor,
)


class WorkbookCloseSpy:
    """记录 excel_sync 每次打开的工作簿，并在实例上挂钩 close 以观察是否真被调用。

    挂钩点是 ``excel_sync.load_catalogue_workbook``（不是 openpyxl 的
    ``load_workbook``）：excel_sync 读基准工作簿一律走这个包装，由它处理空页边距
    兼容副本。``write_approved_plan`` 末尾那个 ``read_only=True`` 的 probe 仍直接
    调 openpyxl，本 spy 覆盖不到它 —— 那条路径由用例末尾真实的 ``os.replace`` 兜底。
    """

    def __init__(self) -> None:
        self.real = excel_sync.load_catalogue_workbook
        self.records: list[dict] = []

    def __call__(self, *args, **kwargs):
        book = self.real(*args, **kwargs)
        record: dict = {"closed": False}
        original = book.close

        def close() -> None:
            record["closed"] = True
            original()

        book.close = close
        self.records.append(record)
        return book

    def install(self):
        patcher = mock.patch.object(excel_sync, "load_catalogue_workbook", self)
        patcher.start()
        return patcher


def make_taxonomy_workbook(path: Path, knowledge_point_id: str = "ZCSQG20260821SF01") -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "目录"
    sheet.cell(1, 1).value = "专题3：整式方程"
    sheet.cell(2, 2).value = "【大题】"
    sheet.cell(3, 3).value = "解一元一次方程"
    sheet.cell(3, 5).value = knowledge_point_id
    book.save(path)
    book.close()


def approved_plan(baseline: Path, **extra) -> dict:
    plan = {"status": "approved", "baseline_sha256": sha256(baseline), "sheet": "目录"}
    plan.update(extra)
    return plan


class WorkbookReleaseTests(unittest.TestCase):
    """openpyxl 的 Workbook 不支持 with；read_only 句柄不 close 会一直锁住文件。"""

    def setUp(self) -> None:
        self.spy = WorkbookCloseSpy()
        self.patcher = self.spy.install()
        self.addCleanup(self.patcher.stop)

    def assert_all_closed(self) -> None:
        self.assertTrue(self.spy.records, "本用例应至少打开过一个工作簿")
        self.assertEqual([record["closed"] for record in self.spy.records], [True] * len(self.spy.records))

    def test_inspect_workbook_closes_workbook_when_topic_missing(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "taxonomy.xlsx"
            make_taxonomy_workbook(path)
            with self.assertRaises(ValueError):
                inspect_workbook(path, "专题99：不存在")
        self.assert_all_closed()

    def test_write_approved_plan_closes_workbook(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            baseline, output = folder / "baseline.xlsx", folder / "output.xlsx"
            make_taxonomy_workbook(baseline)
            plan_path = folder / "plan.json"
            plan_path.write_text(
                json.dumps(approved_plan(baseline, insert_at_row=4, rows=[{"level": 3, "title": "新增三级"}]), ensure_ascii=False),
                encoding="utf-8",
            )
            result = write_approved_plan(baseline, output, plan_path)
            self.assertEqual(result["inserted_rows"], 1)
            self.assert_all_closed()
            # 真实句柄验证：write_approved_plan 末尾的 probe 以 read_only=True 打开 output，
            # 若 probe 未 close，Windows 上这一步会 PermissionError（原 next_sf_number 用例删除后，
            # 这是唯一覆盖 read_only 句柄释放的断言，勿删）。
            os.replace(output, folder / "renamed.xlsx")

    def test_write_approved_refactor_closes_workbook(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            baseline, output = folder / "baseline.xlsx", folder / "output.xlsx"
            make_taxonomy_workbook(baseline)
            plan_path = folder / "plan.json"
            plan_path.write_text(
                json.dumps(
                    approved_plan(
                        baseline,
                        replace_start_row=3,
                        replace_end_row=3,
                        rows=[{"level": 3, "title": "重构后的三级", "knowledge_point_id": "ZCSQG20260917SF09", "question_count": 5}],
                    ),
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            result = write_approved_refactor(baseline, output, plan_path)
            self.assertEqual(result["written_rows"], 1)
            self.assert_all_closed()
            os.replace(output, folder / "renamed.xlsx")

    def test_write_approved_refactor_closes_on_range_error(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            baseline, output = folder / "baseline.xlsx", folder / "output.xlsx"
            make_taxonomy_workbook(baseline)
            plan_path = folder / "plan.json"
            plan_path.write_text(
                json.dumps(
                    approved_plan(baseline, replace_start_row=2, replace_end_row=99, rows=[{"level": 3, "title": "越界"}]),
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            with self.assertRaises(ValueError):
                write_approved_refactor(baseline, output, plan_path)
        self.assert_all_closed()


if __name__ == "__main__":
    unittest.main()
