from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from server.excel_sync import sha256, write_approved_refactor  # noqa: E402


def make_workbook(path: Path) -> None:
    """三行样本：二级容器 + 两个既有三级，P 列有系统目录 ID。"""
    book = Workbook()
    sheet = book.active
    sheet.title = "目录"
    sheet.cell(1, 2).value = "【大题】"
    sheet.cell(2, 3).value = "解分式方程"
    sheet.cell(2, 5).value = "ZCSQG20260821CWJ17"
    sheet.cell(2, 16).value = 760593757
    sheet.cell(3, 3).value = "解不等式组"
    sheet.cell(3, 5).value = "ZCSQG20260821CWJ31"
    sheet.cell(3, 16).value = 760593760
    book.save(path)
    book.close()


def plan_of(baseline: Path, rows: list[dict]) -> dict:
    return {
        "status": "approved",
        "baseline_sha256": sha256(baseline),
        "sheet": "目录",
        "replace_start_row": 2,
        "replace_end_row": 3,
        "rows": rows,
    }


class CatalogIdInheritanceTests(unittest.TestCase):
    """P 列目录 ID 由系统维护：沿用实体必须原值继承，新增实体留空且不得编造。

    见 references/excel-integrity-protocol.md 第 4 节第 5 项。
    """

    def write_plan(self, folder: Path, baseline: Path, rows: list[dict]) -> Path:
        plan_path = folder / "plan.json"
        plan_path.write_text(json.dumps(plan_of(baseline, rows), ensure_ascii=False), encoding="utf-8")
        return plan_path

    def test_inherited_catalog_id_is_written_back(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            baseline, output = folder / "baseline.xlsx", folder / "output.xlsx"
            make_workbook(baseline)
            plan_path = self.write_plan(folder, baseline, [
                {"level": 3, "title": "解分式方程", "catalog_id": 760593757, "question_count": 207},
                {"level": 4, "title": "考法1：解分式方程", "knowledge_point_id": "ZCSQG20260821CWJ17"},
                {"level": 3, "title": "解不等式组", "catalog_id": "760593760", "question_count": 814},
            ])
            write_approved_refactor(baseline, output, plan_path)
            book = load_workbook(output)
            try:
                sheet = book["目录"]
                self.assertEqual(760593757, sheet.cell(2, 16).value)
                # 新增实体：P 列必须留空，等系统导入时生成。
                self.assertIsNone(sheet.cell(3, 16).value)
                # 纯数字串按数字写回，别让它变成文本 ID。
                self.assertEqual(760593760, sheet.cell(4, 16).value)
            finally:
                book.close()

    def test_catalog_id_absent_from_baseline_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            baseline, output = folder / "baseline.xlsx", folder / "output.xlsx"
            make_workbook(baseline)
            plan_path = self.write_plan(folder, baseline, [
                {"level": 3, "title": "解分式方程", "catalog_id": 999999999},
            ])
            with self.assertRaises(ValueError) as caught:
                write_approved_refactor(baseline, output, plan_path)
            self.assertIn("不存在", str(caught.exception))
            self.assertFalse(output.exists())

    def test_duplicate_catalog_ids_are_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            baseline, output = folder / "baseline.xlsx", folder / "output.xlsx"
            make_workbook(baseline)
            plan_path = self.write_plan(folder, baseline, [
                {"level": 3, "title": "甲", "catalog_id": 760593757},
                {"level": 3, "title": "乙", "catalog_id": 760593757},
            ])
            with self.assertRaises(ValueError) as caught:
                write_approved_refactor(baseline, output, plan_path)
            self.assertIn("重复", str(caught.exception))
            self.assertFalse(output.exists())

    def test_new_row_code_colliding_outside_range_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            baseline, output = folder / "baseline.xlsx", folder / "output.xlsx"
            make_workbook(baseline)
            # 只替换第 2 行，却把第 3 行已占用的编号写给新行 ⇒ 必须拦下。
            plan_path = folder / "plan.json"
            plan = plan_of(baseline, [
                {"level": 3, "title": "解分式方程", "knowledge_point_id": "ZCSQG20260821CWJ31"},
            ])
            plan["replace_end_row"] = 2
            plan_path.write_text(json.dumps(plan, ensure_ascii=False), encoding="utf-8")
            with self.assertRaises(ValueError) as caught:
                write_approved_refactor(baseline, output, plan_path)
            self.assertIn("撞号", str(caught.exception))


if __name__ == "__main__":
    unittest.main()
