from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

import yaml
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[2]

from server.main import ServiceState
from server.taxonomy_sync import TaxonomySyncError, discover_latest_workbook, synchronize_taxonomy


def _write_workbook(path: Path, level3_title: str) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "目录"
    sheet.cell(1, 1).value = "专题1：实数"
    sheet.cell(2, 2).value = "【大题】"
    sheet.cell(3, 3).value = level3_title
    book.save(path)
    book.close()


class TaxonomySyncTests(unittest.TestCase):
    def test_selects_latest_date_then_revision_and_rebuilds_only_when_changed(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            old_workbook = root / "目录,ID-3777-2026-09-05 v5.xlsx"
            newest_workbook = root / "目录,ID-3777-2026-09-12 v1.xlsx"
            _write_workbook(old_workbook, "旧目录")
            _write_workbook(newest_workbook, "新目录")
            settings_path = root / "settings.local.yaml"
            output_path = root / "taxonomy.yaml"

            self.assertEqual(discover_latest_workbook(old_workbook), newest_workbook.resolve())
            first = synchronize_taxonomy(
                config_path=settings_path,
                workbook_settings={"path": str(old_workbook), "sheet": "目录"},
                output_path=output_path,
            )
            second = synchronize_taxonomy(
                config_path=settings_path,
                workbook_settings={"path": str(old_workbook), "sheet": "目录"},
                output_path=output_path,
            )

            self.assertEqual(first.status, "updated")
            self.assertEqual(first.workbook_path, newest_workbook.resolve())
            self.assertEqual(second.status, "up_to_date")
            self.assertIn("新目录", output_path.read_text(encoding="utf-8"))

    def test_rejects_missing_unversioned_workbook(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            with self.assertRaises(TaxonomySyncError):
                discover_latest_workbook(Path(directory) / "目录.xlsx")

    def test_rebuilds_when_export_schema_changes_even_if_workbook_is_unchanged(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            workbook = root / "目录,ID-3777-2026-09-12 v1.xlsx"
            _write_workbook(workbook, "实数计算")
            settings_path = root / "settings.local.yaml"
            output_path = root / "taxonomy.yaml"
            synchronize_taxonomy(
                config_path=settings_path,
                workbook_settings={"path": str(workbook), "sheet": "目录"},
                output_path=output_path,
            )
            previous = yaml.safe_load(output_path.read_text(encoding="utf-8"))
            previous["taxonomy_version"] = "excel-v2-stale-export"
            output_path.write_text(yaml.safe_dump(previous, allow_unicode=True, sort_keys=False), encoding="utf-8")

            result = synchronize_taxonomy(
                config_path=settings_path,
                workbook_settings={"path": str(workbook), "sheet": "目录"},
                output_path=output_path,
            )

            self.assertEqual(result.status, "updated")
            self.assertIn("excel-v4-", output_path.read_text(encoding="utf-8"))

    def test_running_service_switches_to_new_workbook_without_restart(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            first_workbook = root / "目录,ID-3777-2026-09-12 v1.xlsx"
            next_workbook = root / "目录,ID-3777-2026-09-16 v1.xlsx"
            _write_workbook(first_workbook, "旧目录")
            settings_path = root / "settings.yaml"
            settings_path.write_text(yaml.safe_dump({
                "host": "127.0.0.1", "port": 0,
                "taxonomy_path": "taxonomy.yaml",
                "rules_path": str(ROOT / "config" / "classification-rules.yaml"),
                "cache_path": "cache.sqlite3",
                "directory_workbook": {"path": str(first_workbook), "sheet": "目录"},
            }, allow_unicode=True), encoding="utf-8")
            state = ServiceState(settings_path)
            try:
                old_version = state.taxonomy.version
                self.assertIn("旧目录", str(state.taxonomy.raw))
                _write_workbook(next_workbook, "新目录")

                self.assertTrue(state.refresh_taxonomy())
                self.assertNotEqual(state.taxonomy.version, old_version)
                self.assertIn("新目录", str(state.taxonomy.raw))
            finally:
                state.close()


if __name__ == "__main__":
    unittest.main()
