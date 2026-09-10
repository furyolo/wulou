"""从目录工作簿只读导出题湖分类所需的专题 / 大题目录 YAML。"""

from __future__ import annotations

import argparse
import hashlib
import re
from pathlib import Path
from typing import Any

import yaml
from openpyxl import load_workbook


# Excel 目录约定：A 列是一级专题，B/C/D 分别是二/三/四级目录。
TOPIC_PATTERN = re.compile(r"^专题\s*(\d+)\s*[：:]?\s*(.+)$")
EXPORT_SCHEMA_VERSION = "v2"


def node_id(prefix: str, row: int, title: str) -> str:
    digest = hashlib.sha1(title.encode("utf-8")).hexdigest()[:10]
    return f"{prefix}-{row}-{digest}"


def _text(value: Any) -> str:
    return str(value or "").strip()


def export_taxonomy(workbook_path: Path, sheet_name: str) -> dict[str, Any]:
    # 使用普通读取模式取得可靠的 max_row；函数从不保存工作簿。
    book = load_workbook(workbook_path, read_only=False, data_only=False)
    try:
        sheet = book[sheet_name]
        topic_rows: list[tuple[int, str, int]] = []
        for row in range(1, sheet.max_row + 1):
            title = _text(sheet.cell(row, 1).value)
            match = TOPIC_PATTERN.match(title)
            if match:
                topic_rows.append((row, title, int(match.group(1))))

        topics: list[dict[str, Any]] = []
        for index, (topic_row, topic_title, topic_order) in enumerate(topic_rows):
            end = topic_rows[index + 1][0] if index + 1 < len(topic_rows) else sheet.max_row + 1
            large_row = next(
                (row for row in range(topic_row + 1, end) if _text(sheet.cell(row, 2).value) == "【大题】"),
                None,
            )
            if large_row is None:
                continue

            # B 列出现下一个二级目录时，【大题】范围结束；不得继承前一个 B 列标题。
            stop = next(
                (row for row in range(large_row + 1, end) if _text(sheet.cell(row, 2).value)),
                end,
            )
            level3: list[dict[str, Any]] = []
            active: dict[str, Any] | None = None
            for row in range(large_row + 1, stop):
                level3_title = _text(sheet.cell(row, 3).value)
                level4_title = _text(sheet.cell(row, 4).value)
                knowledge_point_id = _text(sheet.cell(row, 5).value)
                if level3_title:
                    active = {
                        "id": node_id("l3", row, level3_title),
                        "title": level3_title,
                        "knowledge_point_id": knowledge_point_id or None,
                        "level4": [],
                    }
                    level3.append(active)
                elif level4_title and active is not None:
                    active["level4"].append({
                        "id": node_id("l4", row, level4_title),
                        "title": level4_title,
                        "knowledge_point_id": knowledge_point_id or None,
                    })

            if level3:
                topics.append({
                    "id": node_id("topic", topic_row, topic_title),
                    "title": topic_title,
                    "order": topic_order,
                    "level2": [{
                        "id": node_id("large", large_row, topic_title),
                        "title": "【大题】",
                        "level3": level3,
                    }],
                })
    finally:
        book.close()

    if not topics:
        raise ValueError("未从工作簿 A 列专题范围内提取到任何【大题】目录")
    digest = hashlib.sha256(workbook_path.read_bytes()).hexdigest()[:12]
    return {
        "taxonomy_version": f"excel-{EXPORT_SCHEMA_VERSION}-{digest}",
        "source_workbook": str(workbook_path.resolve()),
        "topics": topics,
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--sheet", default="目录")
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    data = export_taxonomy(args.workbook, args.sheet)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("w", encoding="utf-8", newline="\n") as file:
        yaml.safe_dump(data, file, allow_unicode=True, sort_keys=False)
    print(f"已导出 {len(data['topics'])} 个专题：{args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
