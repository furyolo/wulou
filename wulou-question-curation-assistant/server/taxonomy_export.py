"""从已确认的 Excel 工作簿导出分类目录快照。"""

from __future__ import annotations

import hashlib
import io
import logging
import re
import tempfile
from pathlib import Path
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile

import yaml
from openpyxl import load_workbook


TOPIC_PATTERN = re.compile(r"^专题\s*(\d+)\s*[：:]?\s*(.+)$")
EXPORT_SCHEMA_VERSION = "v4"
LOGGER = logging.getLogger(__name__)
_EMPTY_PAGE_MARGIN = re.compile(rb'\s(?:left|right|top|bottom)=""')
_AUDIT_EVIDENCE_SUFFIX = re.compile(
    r"^(?P<basis>.*?)(?:\s*[；;]\s*|\s+)(?:(?:已有\s*|现有\s*|审计\s*)?证据(?:题目)?\s*(?:ID|编号|题号)|"
    r"evidence\s+(?:exercise|question)\s*(?:ids?|numbers?))\s*[：:]\s*(?P<evidence>.+?)\s*$",
    re.IGNORECASE,
)
_AUDIT_EVIDENCE_ONLY = re.compile(
    r"^(?:(?:已有\s*|现有\s*|审计\s*)?证据(?:题目)?\s*(?:ID|编号|题号)|evidence\s+(?:exercise|question)\s*(?:ids?|numbers?))"
    r"\s*[：:]\s*(?P<evidence>.+?)\s*$",
    re.IGNORECASE,
)
_EVIDENCE_IDENTIFIER = re.compile(r"[A-Za-z][A-Za-z0-9_-]{2,}|\d{3,}")


def node_id(prefix: str, row: int, title: str) -> str:
    digest = hashlib.sha1(title.encode("utf-8")).hexdigest()[:10]
    return f"{prefix}-{row}-{digest}"


def _text(value: Any) -> str:
    return str(value or "").strip()


def split_classification_basis(value: Any) -> tuple[str | None, list[str]]:
    """将目录语义与题号证据分开，题号只供审计而不进入模型目录。"""
    source = _text(value)
    if not source:
        return None, []
    match = _AUDIT_EVIDENCE_SUFFIX.match(source)
    if match:
        basis = match.group("basis").strip(" \t；;") or None
        evidence_text = match.group("evidence")
    else:
        only_match = _AUDIT_EVIDENCE_ONLY.match(source)
        if not only_match:
            return source, []
        basis = None
        evidence_text = only_match.group("evidence")
    return basis, list(dict.fromkeys(_EVIDENCE_IDENTIFIER.findall(evidence_text)))


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as file:
        for chunk in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _compatible_workbook_copy(workbook_path: Path) -> Path | None:
    """修复部分 Excel 导出器写出的空白页边距，仅用于内存外的临时读取副本。"""
    changed = False
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as file:
        temporary_path = Path(file.name)
    try:
        with ZipFile(workbook_path) as source, ZipFile(temporary_path, "w", ZIP_DEFLATED) as destination:
            for member in source.infolist():
                content = source.read(member.filename)
                if member.filename.startswith("xl/worksheets/") and member.filename.endswith(".xml"):
                    normalized = _EMPTY_PAGE_MARGIN.sub(b"", content)
                    changed = changed or normalized != content
                    content = normalized
                destination.writestr(member, content)
    except Exception:
        temporary_path.unlink(missing_ok=True)
        raise
    if not changed:
        temporary_path.unlink(missing_ok=True)
        return None
    return temporary_path


def load_catalogue_workbook(workbook_path: Path):
    """打开目录工作簿；空页边距导致读不了时改读一份剥掉空边距的临时副本。

    ⚠️ 一律先把字节读进内存再交给 openpyxl，不把手柄递给它。原因有二：

    1. **锁文件。** `load_workbook(路径)` 解析到一半抛异常时，它已经打开的
       文件句柄**不会被关掉**，会在 Windows 上一直占着这个工作簿——表现为
       工作簿改不了名、删不掉、Excel 里打不开。而这条失败路径恰好就是
       "目录工作簿带空边距"的情形，也就是题湖系统导出的常态，躲不开。
       先 `read_bytes()` 再喂 `BytesIO`，磁盘句柄在读完那一刻就关了。
    2. **异常类型。** openpyxl 3.1.5 把 `float('')` 的 `ValueError` 包成
       `TypeError('expected <class 'float'>')` 抛出来，所以历史上只捕
       `TypeError` 也能跑；但那是它的内部包装行为，换版本就可能直接抛
       `ValueError`。两个都捕，才不会哪天服务端启动直接崩。

    捕得宽不等于吞异常：副本里没做任何改动时 `_compatible_workbook_copy`
    返回 ``None``，原异常照旧向上抛。

    返回的就是工作簿本身，不把临时副本交给调用方去清理：``read_only=False``
    的 openpyxl 在 ``read()`` 里一次性读完全部内容，副本的字节早已进内存，
    留在磁盘上再删只是徒增一条"调用方忘了删"的路径。读完即删，只有一处
    负责释放。
    """
    try:
        return load_workbook(io.BytesIO(workbook_path.read_bytes()),
                             read_only=False, data_only=False)
    except (TypeError, ValueError):
        compatible_path = _compatible_workbook_copy(workbook_path)
        if compatible_path is None:
            raise
        try:
            book = load_workbook(io.BytesIO(compatible_path.read_bytes()),
                                 read_only=False, data_only=False)
        finally:
            compatible_path.unlink(missing_ok=True)
        LOGGER.warning("目录工作簿含空白页边距，已使用临时兼容副本读取：%s", workbook_path)
        return book


# 旧名保留：历史上临时诊断脚本按 `_load_workbook` 调用过，改名不再让它报错。
_load_workbook = load_catalogue_workbook


def export_taxonomy(workbook_path: Path, sheet_name: str) -> dict[str, Any]:
    """只读提取 A、B、C、D、E、N 列定义的目录及其分类边界。"""
    source_path = workbook_path.resolve()
    book = load_catalogue_workbook(source_path)
    try:
        sheet = book[sheet_name]
        topic_rows: list[tuple[int, str, int]] = []
        for row in range(1, sheet.max_row + 1):
            title = _text(sheet.cell(row, 1).value)
            match = TOPIC_PATTERN.match(title)
            if match:
                topic_rows.append((row, title, int(match.group(1))))

        topics: list[dict[str, Any]] = []
        for index, (topic_row, topic_title, topic_order) in enumerate(topic_rows):
            end = topic_rows[index + 1][0] if index + 1 < len(topic_rows) else sheet.max_row + 1
            large_row = next(
                (row for row in range(topic_row + 1, end) if _text(sheet.cell(row, 2).value) == "【大题】"),
                None,
            )
            if large_row is None:
                continue

            # B 列出现下一个二级目录时，【大题】范围结束；不得继承前一个 B 列标题。
            stop = next(
                (row for row in range(large_row + 1, end) if _text(sheet.cell(row, 2).value)),
                end,
            )
            level3: list[dict[str, Any]] = []
            active: dict[str, Any] | None = None
            for row in range(large_row + 1, stop):
                level3_title = _text(sheet.cell(row, 3).value)
                level4_title = _text(sheet.cell(row, 4).value)
                knowledge_point_id = _text(sheet.cell(row, 5).value)
                classification_basis, audit_evidence_ids = split_classification_basis(sheet.cell(row, 14).value)
                audit_fields = {"audit_evidence_exercise_ids": audit_evidence_ids} if audit_evidence_ids else {}
                if level3_title:
                    active = {
                        "id": node_id("l3", row, level3_title), "source_row": row,
                        "title": level3_title,
                        "knowledge_point_id": knowledge_point_id or None,
                        "classification_basis": classification_basis,
                        "level4": [],
                        **audit_fields,
                    }
                    level3.append(active)
                elif level4_title and active is not None:
                    active["level4"].append({
                        "id": node_id("l4", row, level4_title), "source_row": row,
                        "title": level4_title,
                        "knowledge_point_id": knowledge_point_id or None,
                        "classification_basis": classification_basis,
                        **audit_fields,
                    })

            if level3:
                topics.append({
                    "id": node_id("topic", topic_row, topic_title), "source_row": topic_row,
                    "title": topic_title,
                    "order": topic_order,
                    "level2": [{
                        "id": node_id("large", large_row, topic_title), "source_row": large_row,
                        "title": "【大题】",
                        "level3": level3,
                    }],
                })
    finally:
        book.close()

    if not topics:
        raise ValueError("未从工作簿 A 列专题范围内提取到任何【大题】目录")
    source_sha256 = sha256_file(source_path)
    return {
        "taxonomy_version": f"excel-{EXPORT_SCHEMA_VERSION}-{source_sha256[:12]}",
        "source_workbook": str(source_path),
        "source_sheet": sheet_name,
        "source_workbook_sha256": source_sha256,
        "topics": topics,
    }


def dump_taxonomy(data: dict[str, Any], output_path: Path) -> None:
    """以统一编码和换行写入导出的 YAML。调用方负责原子替换。"""
    with output_path.open("w", encoding="utf-8", newline="\n") as file:
        yaml.safe_dump(data, file, allow_unicode=True, sort_keys=False)
