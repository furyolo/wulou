"""根据与基准工作簿绑定的 taxonomy 自动定位 Focus 的 Excel 范围。"""

from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


class ExcelScopeError(ValueError):
    """基准文件、目录版本或 Focus 锚点不满足安全定位条件。"""


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as file:
        for chunk in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _text(value: Any) -> str:
    return str(value or "").strip()


def _node(items: list[dict[str, Any]], node_id: str, label: str) -> dict[str, Any]:
    item = next((row for row in items if str(row.get("id")) == node_id), None)
    if not item:
        raise ExcelScopeError(f"Focus 的{label}不在当前目录版本中")
    return item


def _next_boundary(sheet: Any, start_row: int, columns: tuple[int, ...]) -> int:
    for row in range(start_row, sheet.max_row + 1):
        if any(_text(sheet.cell(row, column).value) for column in columns):
            return row
    return sheet.max_row + 1


def resolve_excel_scope(settings: dict[str, Any], config_path: Path, taxonomy_raw: dict[str, Any], focus: dict[str, Any]) -> dict[str, Any]:
    """返回自动定位的可替换子目录范围，二级容器行始终不包含在范围内。"""
    configured = settings.get("directory_workbook")
    if not isinstance(configured, dict):
        raise ExcelScopeError("未配置 directory_workbook，不能自动定位 Excel 范围")
    raw_path = str(configured.get("path", "")).strip()
    sheet_name = str(configured.get("sheet", "目录")).strip()
    if not raw_path or not sheet_name:
        raise ExcelScopeError("directory_workbook 必须包含 path 和 sheet")
    baseline = (config_path.parent / raw_path).resolve()
    if not baseline.is_file():
        raise ExcelScopeError("配置的基准工作簿不存在")

    expected_hash = str(taxonomy_raw.get("source_workbook_sha256", "")).strip()
    actual_hash = _sha256(baseline)
    if not expected_hash or actual_hash != expected_hash:
        raise ExcelScopeError("基准工作簿与当前 taxonomy 不一致；请先重新导出目录后再生成方案")
    if str(taxonomy_raw.get("source_sheet", "")).strip() != sheet_name:
        raise ExcelScopeError("配置工作表与当前 taxonomy 来源不一致；请先重新导出目录")

    topic = _node(taxonomy_raw.get("topics") or [], str(focus.get("topic_id", "")), "专题")
    level2 = _node(topic.get("level2") or [], str(focus.get("level2_id", "")), "二级目录")
    level = int(focus.get("level", 0))
    if level not in {2, 3}:
        raise ExcelScopeError("Focus 只能是二级或三级目录")
    level3 = None
    if level == 3:
        level3 = _node(level2.get("level3") or [], str(focus.get("level3_id", "")), "三级目录")

    # 该工作簿在 read_only 模式下可能没有可靠的 max_row，普通打开但绝不保存。
    book = load_workbook(baseline, read_only=False, data_only=False)
    try:
        if sheet_name not in book.sheetnames:
            raise ExcelScopeError("配置的工作表不在基准工作簿中")
        sheet = book[sheet_name]
        level2_row = int(level2.get("source_row", 0))
        if level2_row < 1 or _text(sheet.cell(level2_row, 2).value) != _text(level2.get("title")):
            raise ExcelScopeError("二级目录 Excel 锚点失效；请重新导出目录")
        if level == 2:
            replace_start = level2_row + 1
            replace_end = _next_boundary(sheet, replace_start, (2,)) - 1
            return {
                "baseline_path": str(baseline), "baseline_sha256": actual_hash, "sheet": sheet_name,
                "focus_level": 2, "container_row": level2_row,
                "replace_start_row": replace_start, "replace_end_row": replace_end,
                "preserved_anchor": {"column": "B", "row": level2_row, "title": _text(level2.get("title"))},
            }

        assert level3 is not None
        level3_row = int(level3.get("source_row", 0))
        if level3_row < 1 or _text(sheet.cell(level3_row, 3).value) != _text(level3.get("title")):
            raise ExcelScopeError("三级目录 Excel 锚点失效；请重新导出目录")
        replace_start = level3_row + 1
        replace_end = _next_boundary(sheet, replace_start, (2, 3)) - 1
        return {
            "baseline_path": str(baseline), "baseline_sha256": actual_hash, "sheet": sheet_name,
            "focus_level": 3, "container_row": level3_row,
            "replace_start_row": replace_start, "replace_end_row": replace_end,
            "preserved_anchor": {"column": "C", "row": level3_row, "title": _text(level3.get("title"))},
        }
    finally:
        book.close()
